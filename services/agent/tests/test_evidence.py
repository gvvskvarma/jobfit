import base64
import io
import json
import os
import time
from pathlib import Path

import docx
import psycopg
import pytest
from fastapi.testclient import TestClient

from app.evidence import embed_pending, extract_text, search_evidence, verbatim_only
from app.main import app, signature
from app.schemas import EvidenceDraft

PDF = Path(__file__).with_name("fixtures") / "sample_resume.pdf"

# The sample resume's bullets, tagged the way split_resume tags them.
BANK = {
    "lcp": ("Cut the checkout page's largest contentful paint from 4.1s to 1.6s by "
            "code-splitting the payment widgets and preloading fonts.",
            ["performance", "react", "code-splitting"]),
    "hooks": ("Led the migration of 40 React class components to hooks and TypeScript, "
              "with no regressions in the visual test suite.",
              ["react", "typescript", "migration", "testing"]),
    "sdui": ("Built a server-driven UI layer in GraphQL so merchandisers could change the "
             "home page without a deploy.", ["graphql", "frontend"]),
    "mentor": ("Mentored four junior engineers through weekly code reviews and pairing "
               "sessions.", ["mentoring", "code-review"]),
    "claims": ("Wrote a Node.js service that reconciled insurance claims nightly, processing "
               "2 million rows in under 20 minutes.", ["node.js", "performance", "backend"]),
    "ci": ("Set up GitHub Actions pipelines with Playwright end-to-end tests, cutting release "
           "time from a day to an hour.", ["github actions", "playwright", "ci-cd", "testing"]),
    "aws": ("Moved the patient portal's infrastructure to AWS with Terraform, including RDS "
            "Postgres and CloudFront.", ["aws", "terraform", "postgres", "cloudfront"]),
    "a11y": ("Added WCAG 2.1 AA fixes across the patient portal after an accessibility audit.",
             ["accessibility", "wcag"]),
    "ext": ("Built an open-source Chrome extension that summarizes pull requests, with 3,000 "
            "weekly users.", ["chrome-extension", "open-source"]),
}  # fmt: skip


def test_extracts_pdf_text_without_ligatures() -> None:
    text = extract_text("resume.PDF", PDF.read_bytes())
    assert "Priya Raman" in text
    assert "Certified Developer" in text  # the PDF stores "ﬁ" as one ligature glyph


def test_extracts_docx_paragraphs_and_tables_once() -> None:
    doc = docx.Document()
    doc.add_paragraph("Shipped the search page.")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].merge(row.cells[1]).text = "Northwind Retail"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("r.docx", buf.getvalue())
    assert "Shipped the search page." in text
    assert text.count("Northwind Retail") == 1


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("r.txt", b"hello", ".docx or .pdf"),
        ("r.pdf", b"not a pdf", "couldn't be read"),
        ("r.docx", b"PK\x03\x04junk", "couldn't be read"),
    ],
)
def test_rejects_unreadable_files(name: str, data: bytes, message: str) -> None:
    with pytest.raises(ValueError, match=message):
        extract_text(name, data)


def test_keeps_only_verbatim_items() -> None:
    resume = "Led the migration of 40 React\ncomponents to hooks.\nMentored four engineers."
    drafts = [
        EvidenceDraft(kind="BULLET", text="Led the migration of 40 React components to hooks."),
        EvidenceDraft(kind="BULLET", text="Led the migration of 60 React components to hooks."),
        EvidenceDraft(kind="BULLET", text="Mentored four engineers."),
        EvidenceDraft(kind="BULLET", text="mentored  four engineers."),
    ]
    kept = verbatim_only(resume, drafts)
    assert [d.text for d in kept] == [drafts[0].text, drafts[2].text]


def add_bank(db: psycopg.Connection, user_id: str) -> None:
    db.execute("SELECT set_config('app.user_id', %s, true)", (user_id,))
    db.execute('INSERT INTO "User" (id, email) VALUES (%s, %s)', (user_id, f"{user_id}@test"))
    for key, (text, tags) in BANK.items():
        db.execute(
            """INSERT INTO "Evidence" (id, "userId", kind, text, tags)
               VALUES (%s, %s, 'BULLET', %s, %s)""",
            (f"{user_id}-{key}", user_id, text, tags),
        )


@pytest.mark.parametrize(
    ("query", "expected"),
    [
        ("improved web performance and page load speed", "lcp"),
        ("infrastructure as code on AWS", "aws"),
        ("mentoring and growing other engineers", "mentor"),
        ("end-to-end testing and CI/CD pipelines", "ci"),
        ("migrated a codebase to TypeScript", "hooks"),
    ],
)
def test_search_finds_the_right_bullet(db: psycopg.Connection, query: str, expected: str) -> None:
    with db.transaction():
        add_bank(db, "search-a")
        add_bank(db, "search-b")
        assert embed_pending(db, "search-a") == len(BANK)
        assert embed_pending(db, "search-a") == 0

        hits = search_evidence(db, "search-a", query, top_k=3)
        assert f"search-a-{expected}" in [h.evidence_id for h in hits]
        assert all(h.evidence_id.startswith("search-a-") for h in hits)
        # search-b's rows were never embedded, so they can't be found.
        assert search_evidence(db, "search-b", query) == []
        raise psycopg.Rollback


def test_import_endpoint_requires_a_signature_and_a_resume(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("AGENT_HMAC_SECRET", "test-secret")
    client = TestClient(app)
    body = json.dumps({"filename": "notes.txt", "data": base64.b64encode(b"hi").decode()}).encode()
    ts = str(int(time.time()))
    headers = {"x-jobfit-timestamp": ts, "x-jobfit-user": "u1"}

    bad = client.post(
        "/evidence/import", content=body, headers={**headers, "x-jobfit-signature": "0"}
    )
    assert bad.status_code == 401

    sig = signature(os.environ["AGENT_HMAC_SECRET"], ts, "u1", body)
    res = client.post(
        "/evidence/import", content=body, headers={**headers, "x-jobfit-signature": sig}
    )
    assert res.status_code == 422
    assert ".docx or .pdf" in res.json()["detail"]
