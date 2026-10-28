"""The experience bank (SPEC 6, 10): resume import, local embeddings and search_evidence."""

import io
import os
import re
import unicodedata
from functools import cache

import docx
import psycopg
from fastembed import TextEmbedding
from pydantic import BaseModel
from pypdf import PdfReader

from app import llm
from app.schemas import EvidenceDraft, EvidenceHit

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 40_000
EMBEDDING_MODEL = os.environ.get("EMBEDDING_MODEL", "BAAI/bge-small-en-v1.5")
# BGE models expect this instruction on queries, not on the passages being searched.
QUERY_PREFIX = "Represent this sentence for searching relevant passages: "


class EvidenceDrafts(BaseModel):
    items: list[EvidenceDraft]


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if not name.endswith((".docx", ".pdf")):
        raise ValueError("Upload the resume as a .docx or .pdf file.")
    if len(data) > MAX_RESUME_BYTES:
        raise ValueError("The resume file is larger than 5 MB.")
    try:
        if name.endswith(".docx"):
            doc = docx.Document(io.BytesIO(data))
            cells = (c.text for t in doc.tables for row in t.rows for c in row.cells)
            # Merged table cells repeat their text once per grid cell.
            lines = [p.text for p in doc.paragraphs] + list(dict.fromkeys(cells))
        else:
            lines = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
    except Exception as e:  # each library raises its own zoo of errors on a corrupt file
        raise ValueError("That file couldn't be read. Is it a valid .docx or .pdf?") from e
    # NFKC turns PDF ligatures like "ﬁ" back into plain letters.
    text = unicodedata.normalize("NFKC", "\n".join(lines)).strip()
    if not text:
        raise ValueError("No text found in that file. A scanned PDF has none; try the .docx.")
    if len(text) > MAX_RESUME_CHARS:
        raise ValueError(f"The resume is longer than {MAX_RESUME_CHARS:,} characters.")
    return text


def split_resume(text: str) -> list[EvidenceDraft]:
    safe = re.sub(r"</?\s*resume\s*>", "[tag removed]", text, flags=re.I)
    user = f"<resume>\n{safe}\n</resume>"
    out = llm.recorded(
        "split_resume",
        user,
        EvidenceDrafts,
        lambda: llm._call("split_resume", user, EvidenceDrafts, max_tokens=16_000, timeout=90.0),
    )
    return verbatim_only(text, out.items)


def verbatim_only(resume: str, drafts: list[EvidenceDraft]) -> list[EvidenceDraft]:
    """Drop items the model reworded: the bank holds only what the resume actually says."""
    source = llm._normalize(resume)
    seen: set[str] = set()
    kept = []
    for d in drafts:
        key = llm._normalize(d.text)
        if key and key in source and key not in seen:
            seen.add(key)
            kept.append(d)
    return kept


@cache
def embedder() -> TextEmbedding:
    return TextEmbedding(EMBEDDING_MODEL)


def to_vector(texts: list[str]) -> list[str]:
    """pgvector's text format, so no extra driver adapter is needed."""
    return ["[" + ",".join(f"{x:.6f}" for x in v) + "]" for v in embedder().embed(texts)]


def passage(text: str, tags: list[str]) -> str:
    return f"{text} ({', '.join(tags)})" if tags else text


def as_user(conn: psycopg.Connection, user_id: str) -> None:
    conn.execute("SELECT set_config('app.user_id', %s, true)", (user_id,))


def embed_pending(conn: psycopg.Connection, user_id: str) -> int:
    """Embeds this user's rows with no vector yet; the web app clears it when text changes."""
    with conn.transaction():
        as_user(conn, user_id)
        rows = conn.execute(
            'SELECT id, text, tags FROM "Evidence" WHERE "userId" = %s AND embedding IS NULL',
            (user_id,),
        ).fetchall()
        if not rows:
            return 0
        vectors = to_vector([passage(text, tags) for _, text, tags in rows])
        with conn.cursor() as cur:
            cur.executemany(
                'UPDATE "Evidence" SET embedding = %s::vector WHERE id = %s',
                [(v, row[0]) for v, row in zip(vectors, rows, strict=True)],
            )
    return len(rows)


def search_evidence(
    conn: psycopg.Connection, user_id: str, query: str, top_k: int = 8
) -> list[EvidenceHit]:
    [vector] = to_vector([QUERY_PREFIX + query])
    with conn.transaction():
        as_user(conn, user_id)
        # Hybrid search: the small embedding model misses jargon ("page load speed" vs
        # "largest contentful paint"), keywords miss paraphrase. Reciprocal rank fusion
        # (k=60) merges the two rankings; the keyword side ORs the query's terms.
        rows = conn.execute(
            """WITH q AS (
                 SELECT replace(plainto_tsquery('english', %(q)s)::text, '&', '|') AS t
               ), semantic AS (
                 SELECT id, row_number() OVER (ORDER BY embedding <=> %(v)s::vector) AS r
                 FROM "Evidence" WHERE "userId" = %(u)s AND embedding IS NOT NULL
                 ORDER BY embedding <=> %(v)s::vector LIMIT 50
               ), keyword AS (
                 SELECT id, row_number() OVER (ORDER BY ts_rank(doc, t::tsquery) DESC) AS r
                 FROM (SELECT id, to_tsvector('english', text || ' ' || array_to_string(tags, ' '))
                       AS doc FROM "Evidence"
                       WHERE "userId" = %(u)s AND embedding IS NOT NULL) e, q
                 WHERE t <> '' AND doc @@ t::tsquery
                 LIMIT 50
               )
               SELECT e.id, e.text, sum(1.0 / (60 + x.r)) AS score
               FROM (SELECT * FROM semantic UNION ALL SELECT * FROM keyword) x
               JOIN "Evidence" e USING (id)
               GROUP BY e.id, e.text ORDER BY score DESC LIMIT %(k)s""",
            {"v": vector, "q": query, "u": user_id, "k": top_k},
        ).fetchall()
    return [EvidenceHit(evidence_id=i, text=t, score=round(s, 4)) for i, t, s in rows]
